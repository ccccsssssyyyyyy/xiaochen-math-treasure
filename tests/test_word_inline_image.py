"""Word 导出：图文混排时插图按源文档位置原位写入，而非统一挪到末尾。

覆盖：
- 文字夹图（图在句中被文字包夹）→ 图插入到含前后文字的同一段落
- 图在文末（图后无文字）→ 沿用既有末尾追加逻辑，不污染正文
- 无图题 → 行为与改动前一致
"""
import io
from pathlib import Path

from PIL import Image

from mathbank.word_export_helper import build_word_document
from docx import Document
from docx.oxml.ns import qn


def _make_png(path: Path, color=(10, 120, 200)) -> None:
    Image.new("RGB", (120, 90), color).save(path)


def _build(question: dict, tmp_path: Path):
    _make_png(tmp_path / "q1.png")
    data = [{"question": question, "score": 5}]
    out, _ = build_word_document(
        "测试卷", "高一上", "exam",
        data, include_answers=False,
        show_secret=False, show_notice=False,
        uploads_dir=tmp_path,
    )
    return Document(io.BytesIO(out))


def _paragraphs_with_drawing(doc: Document):
    result = []
    for para in doc.paragraphs:
        if para._p.findall(".//" + qn("w:drawing")):
            result.append(para)
    return result


def test_inline_image_is_placed_within_text_paragraph(tmp_path):
    q = {
        "content": "如图 ![fig](q1.png) 所示，求该集合的元素个数。",
        "image_paths": ["q1.png"],
        "question_type": "single_choice",
    }
    doc = _build(q, tmp_path)
    # 含“如图”的段落里必须同时出现插图（原位写入）
    hit = [p for p in doc.paragraphs if "如图" in p.text and p._p.findall(".//" + qn("w:drawing"))]
    assert hit, "夹图未被原位写入含文字的段落"
    # 整篇只有 1 张图，且就在该段
    assert len(doc.inline_shapes) == 1


def test_image_at_bottom_append_when_align_bottom_right(tmp_path):
    # 图在文末 + figure_align=bottom_right → 走末尾独立段追加路径，不污染题干文字段
    q = {
        "content": "求该集合的元素个数。![fig](q1.png)",
        "image_paths": ["q1.png"],
        "question_type": "single_choice",
        "figure_align": "bottom_right",
    }
    doc = _build(q, tmp_path)
    # 题干文字段不含图
    text_with_drawing = [p for p in doc.paragraphs if "元素个数" in p.text and p._p.findall(".//" + qn("w:drawing"))]
    assert not text_with_drawing, "末尾追加图不应污染题干文字段"
    # 但文档里确实有一张图（在末尾独立段落）
    assert len(doc.inline_shapes) == 1


def test_single_image_at_end_defaults_to_floating_right(tmp_path):
    # 单图在文末、未指定 align（默认 right）→ 沿用既有右侧文字环绕浮动，图随文字段
    q = {
        "content": "求该集合的元素个数。![fig](q1.png)",
        "image_paths": ["q1.png"],
        "question_type": "single_choice",
    }
    doc = _build(q, tmp_path)
    # 浮动图以锚定（wp:anchor）方式嵌入，不计入 inline_shapes，但整篇确有 1 张图
    drawings = doc.element.findall(".//" + qn("w:drawing"))
    assert len(drawings) == 1


def test_no_image_plain_question(tmp_path):
    q = {"content": "求该集合的元素个数。", "question_type": "single_choice"}
    doc = _build(q, tmp_path)
    assert len(doc.inline_shapes) == 0
    assert any("元素个数" in p.text for p in doc.paragraphs)
