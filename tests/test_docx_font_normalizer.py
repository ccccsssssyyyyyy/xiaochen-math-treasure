# tests/test_docx_font_normalizer.py
# -*- coding: utf-8 -*-
"""docx 字体归一化模块的契约测试。

锁死核心不变式：
1. 全部 ``<w:rFonts>``（含 docDefaults / style / run 三个层级）都被改成目标字体。
2. theme 属性（asciiTheme 等）必须被清空，避免覆盖字面字段。
3. 文本内容保持不变。
4. 输出仍是合法 docx，能用 python-docx 重新打开。
5. 同一输入归一化两次字节级相同（幂等）。
"""

from __future__ import annotations

import io
import zipfile
from collections import Counter

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from mathbank.docx_font_normalizer import (
    DEFAULT_CJK_FONT,
    DEFAULT_MATH_FONT,
    normalize_docx_fonts,
)


def _build_mixed_fonts_docx() -> bytes:
    """构造含 docDefaults + style + run 三层字体来源 + theme 字段的 docx。"""
    doc = Document()
    style_el = doc.styles["Normal"].element
    rpr = style_el.find(qn("w:rPr"))
    if rpr is None:
        rpr = etree.SubElement(style_el, qn("w:rPr"))
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = etree.SubElement(rpr, qn("w:rFonts"))
    rfonts.set(qn("w:ascii"), "SimSun")
    rfonts.set(qn("w:hAnsi"), "SimSun")
    rfonts.set(qn("w:eastAsia"), "SimSun")
    rfonts.set(qn("w:cs"), "Times New Roman")
    rfonts.set(qn("w:asciiTheme"), "minorHAnsi")

    p1 = doc.add_paragraph("集合 ")
    run = p1.add_run("A = {1, 2, 3}")
    run.font.name = "Times New Roman"
    rpr_run = run._element.get_or_add_rPr()
    rfonts_run = rpr_run.find(qn("w:rFonts"))
    if rfonts_run is None:
        rfonts_run = etree.SubElement(rpr_run, qn("w:rFonts"))
    rfonts_run.set(qn("w:eastAsia"), "新宋体")
    rfonts_run.set(qn("w:ascii"), "Times New Roman")
    rfonts_run.set(qn("w:hAnsiTheme"), "majorHAnsi")

    p2 = doc.add_paragraph("命题∀x∈R")
    for r in p2.runs:
        r.font.name = "宋体"

    doc.add_paragraph("充要条件")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _collect_rfonts(docx_bytes: bytes):
    out = []
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        for part in ("word/document.xml", "word/styles.xml"):
            try:
                xml = z.read(part)
            except KeyError:
                continue
            tree = etree.parse(io.BytesIO(xml))
            for rfonts in tree.iter(qn("w:rFonts")):
                attrs = {k.split("}", 1)[1]: v for k, v in rfonts.attrib.items()}
                out.append((part, attrs))
    return out


def test_normalize_default_target_is_arial_unicode_ms():
    """回归锁死默认字体：必须为 macOS 唯一同时覆盖中文+数学符号的 Arial Unicode MS。
    防止再次被改成不存在的 STIX Two Math。"""
    assert DEFAULT_CJK_FONT == "Arial Unicode MS"
    assert DEFAULT_MATH_FONT == "Arial Unicode MS"


def test_normalize_overwrites_all_rfonts_with_targets():
    raw = _build_mixed_fonts_docx()
    out = normalize_docx_fonts(raw)
    rfonts = _collect_rfonts(out)
    assert rfonts, "归一化后至少应保留一份 <w:rFonts>"

    for _loc, attrs in rfonts:
        for theme_attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            assert theme_attr not in attrs, f"残留 theme 属性 {theme_attr}={attrs.get(theme_attr)}"
        assert attrs.get("ascii") == DEFAULT_CJK_FONT, attrs
        assert attrs.get("hAnsi") == DEFAULT_CJK_FONT, attrs
        assert attrs.get("eastAsia") == DEFAULT_CJK_FONT, attrs
        assert attrs.get("cs") == DEFAULT_MATH_FONT, attrs


def test_normalize_preserves_text_content():
    raw = _build_mixed_fonts_docx()
    out = normalize_docx_fonts(raw)
    doc = Document(io.BytesIO(out))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "集合" in full_text
    assert "A = {1, 2, 3}" in full_text
    assert "命题∀x∈R" in full_text
    assert "充要条件" in full_text


def test_normalize_output_is_valid_docx():
    raw = _build_mixed_fonts_docx()
    out = normalize_docx_fonts(raw)
    doc = Document(io.BytesIO(out))
    assert doc.styles is not None
    assert len(doc.paragraphs) >= 3


def test_normalize_accepts_custom_fonts():
    raw = _build_mixed_fonts_docx()
    out = normalize_docx_fonts(raw, cjk_font="STHeiti Medium", math_font="Arial Unicode MS")
    rfonts = _collect_rfonts(out)
    for _loc, attrs in rfonts:
        assert attrs.get("ascii") == "STHeiti Medium"
        assert attrs.get("hAnsi") == "STHeiti Medium"
        assert attrs.get("eastAsia") == "STHeiti Medium"
        assert attrs.get("cs") == "Arial Unicode MS"


def test_normalize_idempotent():
    raw = _build_mixed_fonts_docx()
    once = normalize_docx_fonts(raw)
    twice = normalize_docx_fonts(once)
    assert once == twice


def test_normalize_removes_stray_theme_attrs_everywhere():
    raw = _build_mixed_fonts_docx()
    out = normalize_docx_fonts(raw)
    rfonts = _collect_rfonts(out)
    assert len(rfonts) >= 2
    theme_counter: Counter[str] = Counter()
    for _loc, attrs in rfonts:
        for theme_attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            if theme_attr in attrs:
                theme_counter[theme_attr] += 1
    assert sum(theme_counter.values()) == 0, f"theme 字段残留: {theme_counter}"