"""Word 导出：LaTeX tabular 生成真正的 Word 表格（含合并单元格）。"""
import io

from docx import Document

from mathbank.word_export_helper import build_word_document


def _build(content: str, tmp_path):
    data = [{"question": {"content": content, "question_type": "fill_in_blank"}, "score": 5}]
    out, _ = build_word_document(
        "测试卷", "高一上", "exam",
        data, include_answers=False,
        show_secret=False, show_notice=False,
        uploads_dir=tmp_path,
    )
    return Document(io.BytesIO(out))


def test_plain_tabular_three_columns():
    content = (
        "完成下列表格。\n"
        "\\begin{tabular}{ccc}\n"
        "\\toprule\n"
        "甲 & 乙 & 丙 \\\\\n"
        "\\midrule\n"
        "1 & 2 & 3 \\\\\n"
        "4 & 5 & 6 \\\\\n"
        "\\bottomrule\n"
        "\\end{tabular}"
    )
    doc = _build(content, None)
    assert len(doc.tables) == 1
    t = doc.tables[0]
    assert len(t.columns) == 3
    assert len(t.rows) == 3
    # 去分隔符后关键内容齐备
    assert "甲" in t.rows[0].cells[0].text
    assert "6" in t.rows[2].cells[2].text


def test_multicolumn_and_multirow_merge():
    content = (
        "完成下列表格。\n"
        "\\begin{tabular}{ccc}\n"
        "\\toprule\n"
        "\\multicolumn{2}{c}{合并列} & 普通 \\\\\n"
        "\\midrule\n"
        "\\multirow{2}{*}{跨行} & a & b \\\\\n"
        " & c & d \\\\\n"
        "\\bottomrule\n"
        "\\end{tabular}"
    )
    doc = _build(content, None)
    assert len(doc.tables) == 1
    t = doc.tables[0]
    assert len(t.columns) == 3
    assert len(t.rows) == 3
    # 横向合并：第 0 行前两格合并，文本落在合并后的格中
    assert "合并列" in t.rows[0].cells[0].text
    # 纵向合并：第 1、2 行第 0 格合并，文本落在合并后的格中
    assert "跨行" in t.rows[1].cells[0].text
    # 其余内容无丢失
    assert "普通" in t.rows[0].cells[2].text
    assert "a" in t.rows[1].cells[1].text
    assert "d" in t.rows[2].cells[2].text
